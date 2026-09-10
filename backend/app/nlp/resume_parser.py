import re
import os
import pymupdf
import docx
from typing import Dict, List, Any, Optional
from app.nlp.skills_dataset import TECH_SKILLS_DATA, categorize_skill

class ResumeParser:
    @staticmethod
    def extract_text_from_pdf(file_path: str) -> str:
        """Extract clean plain text from PDF using PyMuPDF."""
        text = ""
        try:
            doc = pymupdf.open(file_path)
            for page in doc:
                text += page.get_text("text") + "\n"
            doc.close()
        except Exception as e:
            print(f"Error reading PDF {file_path}: {e}")
        return text.strip()

    @staticmethod
    def extract_text_from_docx(file_path: str) -> str:
        """Extract text from DOCX document."""
        text = ""
        try:
            doc = docx.Document(file_path)
            for para in doc.paragraphs:
                if para.text.strip():
                    text += para.text.strip() + "\n"
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        text += row_text + "\n"
        except Exception as e:
            print(f"Error reading DOCX {file_path}: {e}")
        return text.strip()

    @classmethod
    def extract_text(cls, file_path: str) -> str:
        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".pdf":
            return cls.extract_text_from_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return cls.extract_text_from_docx(file_path)
        elif ext == ".txt":
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        return ""

    @staticmethod
    def extract_email(text: str) -> Optional[str]:
        email_pattern = r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+'
        matches = re.findall(email_pattern, text)
        if matches:
            # Filter out generic or dummy emails if needed
            for match in matches:
                clean = match.strip().rstrip(".")
                if len(clean) > 5:
                    return clean
        return None

    @staticmethod
    def extract_phone(text: str) -> Optional[str]:
        # Support international and Indian numbers (+91, 10 digits, dashes, parentheses)
        phone_patterns = [
            r'(\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}',
            r'(\+?91[-.\s]?)?[6789]\d{9}',
            r'\+?\d{1,3}[-.\s]?\d{10}'
        ]
        for pattern in phone_patterns:
            matches = re.findall(pattern, text)
            if matches:
                # Find the full matched segment
                for m in re.finditer(pattern, text):
                    num = m.group(0).strip()
                    if len(re.sub(r'\D', '', num)) >= 10:
                        return num
        return None

    @staticmethod
    def extract_name(text: str, email: Optional[str] = None) -> str:
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        # Typically the candidate name is in the first 3-5 lines before headers like 'Experience' or 'Skills'
        blacklist_keywords = ["curriculum", "resume", "cv", "profile", "contact", "summary", "experience", "education", "skills", "phone", "email", "address"]
        for line in lines[:6]:
            if len(line.split()) in [2, 3, 4] and len(line) < 35:
                lower = line.lower()
                if not any(k in lower for k in blacklist_keywords) and not re.search(r'[@\d]', line):
                    return line.title()
        
        # Fallback to email username if nothing clean found
        if email:
            user_part = email.split("@")[0]
            clean_part = re.sub(r'[\._0-9]', ' ', user_part).strip().title()
            if len(clean_part) > 2:
                return clean_part
        return "Candidate Profile"

    @staticmethod
    def extract_skills(text: str) -> List[str]:
        """Match tech skills using word boundaries and aliases."""
        detected = set()
        lowered_text = " " + text.lower() + " "
        # Replace non-alphanumeric except common skill punctuation
        normalized = re.sub(r'[^a-z0-9\+\#\.\s\_\-]', ' ', lowered_text)

        for canonical_name, aliases in TECH_SKILLS_DATA.items():
            for alias in aliases:
                # Use regex with boundary checking
                escaped = re.escape(alias.lower())
                # Handle special characters like C++, C#, .NET
                pattern = rf'(?<![a-zA-Z0-9]){escaped}(?![a-zA-Z0-9])'
                if re.search(pattern, normalized):
                    detected.add(canonical_name)
                    break

        return sorted(list(detected))

    @staticmethod
    def extract_experience_years(text: str) -> float:
        """Estimate years of experience from resume text."""
        exp_patterns = [
            r'(\d+(?:\.\d+)?)\+?\s*(?:years?|yrs?)(?:\s+of)?\s+experience',
            r'experience\s*(?:of)?\s*(\d+(?:\.\d+)?)\+?\s*(?:years?|yrs?)',
            r'total\s+experience\s*:\s*(\d+(?:\.\d+)?)\s*(?:years?|yrs?)'
        ]
        lowered = text.lower()
        for p in exp_patterns:
            matches = re.findall(p, lowered)
            if matches:
                try:
                    val = float(matches[0])
                    if 0 <= val <= 35:
                        return val
                except ValueError:
                    pass

        # Look for date ranges e.g. 2019 - 2024 or 2021 - Present
        year_ranges = re.findall(r'(20[0-2][0-9])\s*(?:-|–|to)\s*(20[0-2][0-9]|present|current)', lowered)
        if year_ranges:
            total_span = 0
            for start_yr, end_yr in year_ranges:
                try:
                    s = int(start_yr)
                    e = 2026 if end_yr in ["present", "current"] else int(end_yr)
                    diff = max(0, e - s)
                    if diff < 15:
                        total_span += diff
                except ValueError:
                    pass
            if total_span > 0:
                return min(round(total_span, 1), 25.0)

        # Default estimated experience
        return 2.5

    @staticmethod
    def extract_education(text: str) -> str:
        lowered = text.lower()
        edu_keywords = [
            "master of science", "master of technology", "m.tech", "m.s.", "m.sc", "mca",
            "bachelor of technology", "bachelor of science", "bachelor of engineering",
            "b.tech", "b.e.", "b.s.", "b.sc", "bca", "ph.d", "phd", "doctorate",
            "computer science", "information technology", "software engineering",
            "data science", "artificial intelligence", "electrical engineering"
        ]
        found = []
        if "ph.d" in lowered or "phd" in lowered:
            found.append("Ph.D / Doctorate")
        if any(m in lowered for m in ["master", "m.tech", "m.s.", "m.sc", "mca"]):
            found.append("Master's Degree")
        if any(b in lowered for b in ["bachelor", "b.tech", "b.e.", "b.s.", "b.sc", "bca", "undergraduate"]):
            found.append("Bachelor's Degree")
        
        if "computer science" in lowered or "information technology" in lowered:
            found.append("Computer Science & Engineering")
        elif "artificial intelligence" in lowered or "data science" in lowered:
            found.append("AI / Data Science")

        if found:
            return " in ".join(found[:2])
        return "B.S. in Computer Science / Technical Degree"

    @classmethod
    def parse_full_resume(cls, file_path: str) -> Dict[str, Any]:
        """Runs the entire end-to-end extraction pipeline on a resume file."""
        text = cls.extract_text(file_path)
        if not text:
            return {
                "raw_text": "",
                "name": "Candidate",
                "email": None,
                "phone": None,
                "skills": [],
                "skills_by_category": {},
                "experience_years": 0.0,
                "education": "Not specified",
                "summary": "Empty or unreadable document.",
                "total_words": 0
            }

        email = cls.extract_email(text)
        phone = cls.extract_phone(text)
        name = cls.extract_name(text, email)
        skills = cls.extract_skills(text)
        exp_years = cls.extract_experience_years(text)
        education = cls.extract_education(text)

        # Categorize detected skills
        skills_by_cat: Dict[str, List[str]] = {}
        for skill in skills:
            cat = categorize_skill(skill)
            skills_by_cat.setdefault(cat, []).append(skill)

        # Extract summary / objective or first paragraph
        lines = [l.strip() for l in text.split("\n") if l.strip() and len(l.strip()) > 30]
        summary = lines[0] if lines else "Experienced software professional with passion for building high impact software."
        if len(summary) > 280:
            summary = summary[:277] + "..."

        return {
            "raw_text": text,
            "name": name,
            "email": email,
            "phone": phone,
            "skills": skills,
            "skills_by_category": skills_by_cat,
            "experience_years": exp_years,
            "education": education,
            "summary": summary,
            "total_words": len(text.split())
        }
